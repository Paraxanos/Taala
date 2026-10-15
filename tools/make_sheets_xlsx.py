"""Turn the frozen per-author TSV sheets into one Excel workbook per author.

Four separate files, never four tabs in one workbook: the only thing observed_authors.tsv
measures is four people spelling independently, and a shared file lets the second author
read the first one's answers.

The row order comes from the existing TSVs rather than being re-shuffled, so the xlsx and
the tsv stay the same sheet and either can be filled.

Deliberately absent: the reference spelling (survey/words.tsv ref_latin) and any worked
romanization example. survey/instructions.md pre-registers that participants see neither,
because either one anchors the spelling they produce. Do not add an example here.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SHEETS = ROOT / "data" / "variants" / "sheets"
AUTHORS = ("A", "B", "C", "D")

BODY = "Arial"
INDIC = "Nirmala UI"      # Arial has no usable Devanagari; Nirmala UI ships with Windows

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
GIVEN_FILL = PatternFill("solid", fgColor="F2F2F2")   # do not edit
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # fill these
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RULES = [
    ("What to do",
     "For each word, type how YOU would normally write it in English letters."),
    ("Your first instinct",
     "There is no right answer. Do not look the word up, and do not think about what is "
     "correct -- we are measuring what people actually type."),
    ("spelling",
     "Required. Your natural spelling, lowercase a-z only. No spaces, digits or punctuation."),
    ("spelling_alt",
     "Optional. Fill it only if a SECOND spelling feels equally natural to you."),
    ("dont_know",
     "Put 1 here if you do not know the word. Leave spelling empty in that case."),
    ("Work alone",
     "Do not discuss the words with the other three, and do not look at anyone else's "
     "sheet. If two of us agree a spelling first, the variation we are measuring is gone."),
    ("Do not open survey/words.tsv",
     "It holds a reference spelling for every word. Reading it would anchor yours to that one."),
    ("Partial is fine",
     "Go down the list in order and stop when you stop. Sixty words from you beats zero."),
]


def shown(path):
    """A tidy path for printing, without assuming the file sits inside the repo."""
    return path.relative_to(ROOT) if path.is_relative_to(ROOT) else path


def read_sheet(letter):
    with (SHEETS / f"author_{letter}.tsv").open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh, delimiter="\t"))


def write_instructions(ws, letter, n_words):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 86

    ws["A1"] = f"Taala spelling sheet - author {letter}"
    ws["A1"].font = Font(name=BODY, size=16, bold=True, color="1F3864")
    ws["A2"] = "Hindi words in Devanagari. Write how you would type each one in English letters."
    ws["A2"].font = Font(name=BODY, size=11, color="595959")

    row = 4
    for label, text in RULES:
        ws.cell(row, 1, label).font = Font(name=BODY, size=11, bold=True)
        cell = ws.cell(row, 2, text)
        cell.font = Font(name=BODY, size=11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 30
        row += 1

    row += 1
    ws.cell(row, 1, "Columns to fill").font = Font(name=BODY, size=11, bold=True, color="BF8F00")
    ws.cell(row, 2, "The shaded (cream) columns on the Spellings tab: spelling, spelling_alt, "
                    "dont_know. The grey columns are given -- leave them alone.") \
        .font = Font(name=BODY, size=11)
    ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 30

    row += 2
    ws.cell(row, 1, "Words filled in").font = Font(name=BODY, size=11, bold=True)
    done = ws.cell(row, 2, f"=COUNTA(Spellings!C2:C{n_words + 1})&\" of {n_words}\"")
    done.font = Font(name=BODY, size=11, bold=True, color="1F3864")

    row += 2
    ws.cell(row, 1, "When you are done").font = Font(name=BODY, size=11, bold=True)
    ws.cell(row, 2, "Save the file where it is and tell the data dev. Do not rename it -- "
                    "the importer finds it by name.").font = Font(name=BODY, size=11)


def write_spellings(ws, rows):
    headers = ["word_id", "word", "spelling", "spelling_alt", "dont_know"]
    widths = [10, 20, 20, 20, 12]

    for col, (name, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, col, name)
        cell.font = Font(name=BODY, size=11, bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BOX
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22

    for r, row in enumerate(rows, start=2):
        given = [row["word_id"], row["devanagari"]]
        for col, value in enumerate(given, start=1):
            cell = ws.cell(r, col, value)
            cell.fill = GIVEN_FILL
            cell.border = BOX
            if col == 2:
                cell.font = Font(name=INDIC, size=16)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.font = Font(name=BODY, size=10, color="808080")
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in (3, 4, 5):
            cell = ws.cell(r, col)
            cell.fill = INPUT_FILL
            cell.border = BOX
            cell.font = Font(name=BODY, size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 24

    last = len(rows) + 1
    # Advisory only: Excel cannot express "a-z only" without a formula nobody can read, so
    # the shape is checked here and tools/import_sheets.py is what actually enforces it.
    spelling_rule = DataValidation(
        type="custom",
        formula1="=AND(LEN(C2)>=2,LEN(C2)<=24,EXACT(C2,LOWER(C2)),ISERROR(SEARCH(\" \",C2)))",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Check the spelling",
        error="Lowercase a-z only, 2 to 24 letters, no spaces, digits or punctuation.")
    ws.add_data_validation(spelling_rule)
    spelling_rule.add(f"C2:D{last}")

    flag = DataValidation(type="list", formula1='"1"', allow_blank=True, showErrorMessage=True,
                          errorTitle="dont_know", error="Put 1 here, or leave it empty.")
    ws.add_data_validation(flag)
    flag.add(f"E2:E{last}")

    ws.freeze_panes = "C2"


def main():
    for letter in AUTHORS:
        rows = read_sheet(letter)
        wb = Workbook()
        wb.active.title = "Instructions"
        write_instructions(wb.active, letter, len(rows))
        write_spellings(wb.create_sheet("Spellings"), rows)
        out = SHEETS / f"author_{letter}.xlsx"
        wb.save(out)
        print(f"wrote {shown(out)}  ({len(rows)} words)")


if __name__ == "__main__":
    main()
