"""Read the filled author sheets into data/variants/observed_authors.tsv.

Takes author_<L>.xlsx if it has anything in it, otherwise author_<L>.tsv, so either way of
filling the sheet works. Excel's own validation is advisory; this is where the rules are
actually enforced, and every rejected cell is printed rather than dropped quietly.
"""

import csv
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SHEETS = ROOT / "data" / "variants" / "sheets"
WORDS = ROOT / "survey" / "words.tsv"
OUT = ROOT / "data" / "variants" / "observed_authors.tsv"
AUTHORS = ("A", "B", "C", "D")

def shown(path):
    """A tidy path for printing, without assuming the file sits inside the repo."""
    return path.relative_to(ROOT) if path.is_relative_to(ROOT) else path


SPELLING = re.compile(r"^[a-z]{2,24}$")


def load_words():
    with WORDS.open(encoding="utf-8", newline="") as fh:
        return {r["id"]: r for r in csv.DictReader(fh, delimiter="\t")}


def rows_from_tsv(path):
    with path.open(encoding="utf-8", newline="") as fh:
        for r in csv.DictReader(fh, delimiter="\t"):
            yield r["word_id"], r.get("spelling", ""), r.get("spelling_alt", ""), r.get("dont_know", "")


def rows_from_xlsx(path):
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True)["Spellings"]
    for r in range(2, ws.max_row + 1):
        cells = [ws.cell(r, c).value for c in range(1, 6)]
        if cells[0]:
            yield tuple(str(c).strip() if c is not None else "" for c in cells[:1] + cells[2:5])


def read_author(letter):
    """Prefer whichever file the author actually filled in."""
    xlsx, tsv = SHEETS / f"author_{letter}.xlsx", SHEETS / f"author_{letter}.tsv"
    if xlsx.exists():
        rows = list(rows_from_xlsx(xlsx))
        if any(s or a or d for _, s, a, d in rows):
            return rows, xlsx
    if tsv.exists():
        return list(rows_from_tsv(tsv)), tsv
    return [], None


def main():
    words = load_words()
    out_rows, problems, summary = [], [], []

    for letter in AUTHORS:
        rows, source = read_author(letter)
        kept = unknown = blank = 0
        for word_id, spelling, alt, flag in rows:
            if word_id not in words:
                problems.append(f"{letter}: unknown word id {word_id!r}")
                continue
            flag = flag.strip()
            if flag and flag != "1":
                problems.append(f"{letter}/{word_id}: dont_know is {flag!r}, expected 1 or empty")
                continue
            if flag == "1":
                if spelling:
                    problems.append(f"{letter}/{word_id}: marked dont_know but also spelled {spelling!r}")
                unknown += 1
                continue
            if not spelling and not alt:
                blank += 1
                continue
            for kind, value in (("natural", spelling), ("alt", alt)):
                value = value.strip().lower()
                if not value:
                    continue
                if not SPELLING.match(value):
                    problems.append(f"{letter}/{word_id}: {value!r} is not 2-24 lowercase letters")
                    continue
                out_rows.append([word_id, words[word_id]["devanagari"], letter, value, kind])
                kept += 1
        name = source.name if source else "no sheet"
        summary.append(f"  author {letter}: {kept} spellings, {unknown} don't-know, "
                       f"{blank} blank  ({name})")

    seen = set()
    deduped = []
    for row in out_rows:
        key = (row[0], row[2], row[3])
        if key not in seen:
            seen.add(key)
            deduped.append(row)

    with OUT.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter="\t", lineterminator="\n")
        writer.writerow(["word_id", "devanagari", "author", "spelling", "kind"])
        writer.writerows(deduped)

    print("\n".join(summary))
    for p in problems:
        print(f"  REJECTED {p}", file=sys.stderr)
    authors = len({r[2] for r in deduped})
    print(f"wrote {shown(OUT)}: {len(deduped)} spellings from {authors} author(s), "
          f"{len({r[0] for r in deduped})} words covered")
    if problems:
        print(f"{len(problems)} cell(s) rejected -- fix them in the sheet and re-run.")
    if not deduped:
        print("Nothing filled in yet. This is expected until the sheets come back.")


if __name__ == "__main__":
    main()
